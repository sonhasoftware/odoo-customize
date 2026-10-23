"""Validate and extract a single file attached to a chat message."""

import base64
import binascii
import csv
import io
from pathlib import PurePath


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 30000
ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.xls', '.csv', '.txt', '.png', '.jpg', '.jpeg', '.webp'}
IMAGE_MIME = {'.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.webp': 'image/webp'}


class ChatAttachmentError(ValueError):
    pass


def _decode_text(data):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ChatAttachmentError('Tệp văn bản không dùng bảng mã được hỗ trợ (UTF-8 hoặc CP1252).')


def _extract_pdf(env, data, question):
    try:
        import fitz
        pdf = fitz.open(stream=data, filetype='pdf')
        try:
            if pdf.page_count > 40:
                raise ChatAttachmentError('PDF có quá 40 trang. Hãy tải vào Documents của chủ đề để xử lý.')
            parts = []
            scanned_pages = 0
            for page in pdf:
                text = page.get_text().strip()
                if len(text) < 10:
                    scanned_pages += 1
                    if scanned_pages > 5:
                        raise ChatAttachmentError('PDF scan có quá 5 trang ảnh. Hãy tải vào Documents của chủ đề để xử lý.')
                    image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes('png')
                    text = _extract_image(env, image, question)
                parts.append(f'## Trang {page.number + 1}\n{text}')
            return '\n\n'.join(parts)
        finally:
            pdf.close()
    except ChatAttachmentError:
        raise
    except Exception as error:
        raise ChatAttachmentError(f'Không đọc được PDF: {error}') from error


def _extract_docx(data):
    try:
        from docx import Document
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text.strip() for cell in row.cells))
        return '\n'.join(parts)
    except Exception as error:
        raise ChatAttachmentError(f'Không đọc được DOCX: {error}') from error


def _extract_spreadsheet(data, extension):
    try:
        if extension == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                parts = []
                total_chars = 0
                for sheet in workbook:
                    parts.append(f'## Sheet: {sheet.title}')
                    for row in sheet.iter_rows(values_only=True):
                        line = ' | '.join('' if cell is None else str(cell) for cell in row)
                        parts.append(line)
                        total_chars += len(line)
                        if total_chars > MAX_TEXT_CHARS:
                            break
                    if total_chars > MAX_TEXT_CHARS:
                        break
                return '\n'.join(parts)
            finally:
                workbook.close()
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data)
        parts = []
        total_chars = 0
        for sheet in workbook.sheets():
            parts.append(f'## Sheet: {sheet.name}')
            for row_index in range(sheet.nrows):
                line = ' | '.join(str(cell.value) for cell in sheet.row(row_index))
                parts.append(line)
                total_chars += len(line)
                if total_chars > MAX_TEXT_CHARS:
                    break
            if total_chars > MAX_TEXT_CHARS:
                break
        return '\n'.join(parts)
    except Exception as error:
        raise ChatAttachmentError(f'Không đọc được bảng tính: {error}') from error


def _extract_image(env, data, question):
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        from . import image_classifier, ocr_service, vision_service
        category = image_classifier.classify_image(data).get('category', 'mixed')
        vision_prompt = (
            'Hãy mô tả cấu trúc và nội dung nhìn thấy trong ảnh bằng tiếng Việt. '
            'Nếu có chữ hoặc số, chép nguyên văn, không tự sửa hay đoán phần không rõ. '
            'Câu hỏi của người dùng: ' + question
        )
        results = []
        try:
            ocr_text = ocr_service.call_ocr_engine(env, data)
            if ocr_text and not ocr_text.startswith(('[Văn bản trong ảnh:', '[Sơ đồ / Lưu đồ:', '[Lỗi')):
                results.append(f'Văn bản OCR (nguồn ưu tiên cho chữ và số):\n{ocr_text}')
        except Exception:
            pass
        if category != 'text_table' or not results:
            try:
                vision_text = vision_service.call_vision_ocr(env, data, prompt=vision_prompt)
                if vision_text and not vision_text.startswith(('[Văn bản trong ảnh:', '[Sơ đồ / Lưu đồ:', '[Lỗi')):
                    results.append(f'Mô tả ảnh bằng Vision (chỉ bổ trợ cho cấu trúc; nếu chữ/số khác OCR, dùng OCR):\n{vision_text}')
            except Exception:
                pass
        text = '\n\n'.join(results)
        if not text or text.startswith(('[Văn bản trong ảnh:', '[Sơ đồ / Lưu đồ:', '[Lỗi')):
            raise ChatAttachmentError('Không đọc được ảnh bằng OCR/Vision hiện tại. Hãy kiểm tra model và cấu hình OCR/Vision.')
        return text
    except ChatAttachmentError:
        raise
    except Exception as error:
        raise ChatAttachmentError(f'Không phân tích được ảnh: {error}') from error


def prepare_chat_attachment(env, payload, question):
    """Return safe message field values; reject unsupported or unreadable files."""
    if not isinstance(payload, dict):
        raise ChatAttachmentError('Tệp đính kèm không hợp lệ.')
    name = str(payload.get('name') or '').strip()
    if not name or name != PurePath(name).name or '/' in name or '\\' in name or any(ord(char) < 32 for char in name):
        raise ChatAttachmentError('Tên tệp đính kèm không hợp lệ.')
    extension = '.' + name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    if extension not in ALLOWED_EXTENSIONS:
        raise ChatAttachmentError('Chỉ hỗ trợ PDF, DOCX, XLSX, XLS, CSV, TXT, PNG, JPG và WEBP.')
    encoded = payload.get('data')
    if not isinstance(encoded, str) or len(encoded) > (MAX_FILE_BYTES * 4 // 3 + 16):
        raise ChatAttachmentError('Tệp đính kèm vượt quá giới hạn 10 MB.')
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as error:
        raise ChatAttachmentError('Dữ liệu tệp đính kèm không hợp lệ.') from error
    if not data or len(data) > MAX_FILE_BYTES:
        raise ChatAttachmentError('Tệp rỗng hoặc vượt quá giới hạn 10 MB.')

    if extension in IMAGE_MIME:
        text = _extract_image(env, data, question)
        mimetype = IMAGE_MIME[extension]
    elif extension == '.pdf':
        text = _extract_pdf(env, data, question)
        mimetype = 'application/pdf'
    elif extension == '.docx':
        text = _extract_docx(data)
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif extension in ('.xlsx', '.xls'):
        text = _extract_spreadsheet(data, extension)
        mimetype = ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if extension == '.xlsx' else 'application/vnd.ms-excel')
    elif extension == '.csv':
        text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(_decode_text(data))))
        mimetype = 'text/csv'
    else:
        text = _decode_text(data)
        mimetype = 'text/plain'

    text = text.strip()
    if not text:
        raise ChatAttachmentError('Không trích xuất được nội dung từ tệp này.')
    if len(text) > MAX_TEXT_CHARS:
        raise ChatAttachmentError('Nội dung tệp vượt 30.000 ký tự. Hãy tải vào Documents của chủ đề để xử lý.')
    return {
        'attachment_name': name,
        'attachment_mimetype': mimetype,
        'attachment_data': base64.b64encode(data).decode('ascii'),
        'attachment_text': text,
    }
